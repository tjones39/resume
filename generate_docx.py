"""Regenerate Todd_Jones_Resume.docx from index.html (single source of truth)."""
import re, html as H, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor

SRC, OUT = 'index.html', 'Todd_Jones_Resume.docx'
s = open(SRC).read()
s = re.sub(r'<!--.*?-->', '', s, flags=re.S)          # strip comments
s = re.sub(r'<style.*?</style>', '', s, flags=re.S)   # strip CSS

def txt(x):
    # <br> becomes a real break; every other newline is just HTML source
    # wrapping and must collapse to a single space.
    x = re.sub(r'<br\s*/?>', '\x00', x)
    x = re.sub(r'<[^>]+>', '', x)
    x = H.unescape(x)
    x = re.sub(r'\s+', ' ', x)
    return x.replace('\x00', '\n').strip()

def one(pat, label, flags=re.S):
    m = re.search(pat, s, flags)
    if not m: print(f'!! parse fail: {label}'); sys.exit(1)
    return m

name     = txt(one(r'<h1>(.*?)</h1>', 'name').group(1))
subtitle = txt(one(r'class="subtitle">(.*?)</div>', 'subtitle').group(1))
location = txt(one(r'class="location">(.*?)</div>', 'location').group(1))
contact  = [txt(a) for a in re.findall(r'<a[^>]*>(.*?)</a\s*>',
            one(r'class="contact">(.*?)</div>', 'contact').group(1), re.S)]
summary  = txt(one(r'id="summary".*?<p>(.*?)</p>', 'summary').group(1))
skills   = [(txt(a), txt(b)) for a, b in
            re.findall(r'<dt>(.*?)</dt>\s*<dd>(.*?)</dd>',
            one(r'id="skills"(.*?)</dl>', 'skills').group(1), re.S)]
edu      = txt(one(r'class="education">(.*?)</div>', 'education').group(1))

jobs = []
for blk in re.findall(r'<div class="job">(.*?)\n                </div>', s, re.S):
    g = lambda c: (lambda m: txt(m.group(1)) if m else '')(
        re.search(rf'class="{c}">(.*?)</div>', blk, re.S))
    jobs.append({'title': g('job-title'), 'company': g('job-company'),
                 'dates': g('job-dates'), 'loc': g('job-location'),
                 'bullets': [txt(li) for li in re.findall(r'<li>(.*?)</li>', blk, re.S)]})

if not jobs: print('!! no jobs parsed'); sys.exit(1)
print(f'parsed: {len(jobs)} jobs, {len(skills)} skill rows, {len(contact)} contact links')

# ---------- build ----------
BLUE, MID, GREY, SUB = RGBColor(0x19,0x76,0xD2), RGBColor(0x15,0x65,0xC0), \
                       RGBColor(0x88,0x88,0x88), RGBColor(0x55,0x55,0x55)
doc = Document()
n = doc.styles['Normal']; n.font.name = 'Calibri'; n.font.size = Pt(10.5)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.8)
sec.top_margin = sec.bottom_margin = Inches(0.7)

def para(text, *, style=None, bold=None, size=None, color=None,
         before=None, after=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    r = p.add_run(text)
    if bold is not None: r.bold = bold
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    f = p.paragraph_format
    if before is not None: f.space_before = Pt(before)
    if after is not None: f.space_after = Pt(after)
    return p

def heading(t):
    para(t, bold=True, size=13, color=BLUE, before=12, after=2)

para(name, bold=True, size=26, color=BLUE)
para(subtitle, size=13, color=SUB, after=0)
para(location, size=10, color=GREY, after=2)
para('  |  '.join(contact), size=10, color=MID, after=8)

heading('SUMMARY')
para(summary, after=4)

heading('SKILLS')
for label, body in skills:
    p = doc.add_paragraph()
    rl = p.add_run(f'{label}: '); rl.bold = True; rl.font.color.rgb = MID
    rl.font.size = Pt(10.5)
    rb = p.add_run(body); rb.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(2)

heading('EXPERIENCE')
for j in jobs:
    para(j['title'], bold=True, size=11.5, color=BLUE, before=8, after=0)
    meta = '    '.join(x for x in [j['company'], j['dates']] if x)
    if j['loc']: meta += '    |    ' + j['loc']
    para(meta, bold=True, size=10.5, color=MID, after=2)
    for b in j['bullets']:
        para(b, style='List Bullet', size=10.5, after=2)

heading('EDUCATION')
lines = [l.strip() for l in edu.split('\n') if l.strip()]
p = doc.add_paragraph(); r = p.add_run(lines[0]); r.bold = True
for ln in lines[1:]:
    p.add_run('\n' + ln)

doc.save(OUT)
print(f'wrote {OUT}')
